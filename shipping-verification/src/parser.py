import io
from pathlib import Path

# OCR Imports with graceful fallback check
try:
    from pdf2image import convert_from_bytes
    import pytesseract
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False


class DocumentParser:
    """
    Parses plain text, PDF, DOCX, and Excel files into plain text.
    Handles corrupt or unreadable files gracefully and falls back to OCR for scanned PDFs.
    """
    def __init__(self, inbox_instance=None):
        self.inbox = inbox_instance

    def parse(self, attachment_path):
        """
        Reads attachment path and extracts full text content.
        Returns text string or None if unreadable.
        """
        try:
            if self.inbox:
                raw_bytes = self.inbox.read_bytes(attachment_path)
            else:
                raw_bytes = Path(attachment_path).read_bytes()
        except Exception:
            return None

        ext = attachment_path.split(".")[-1].lower() if "." in attachment_path else ""

        if ext in ["txt", "text", "csv", "log"]:
            return raw_bytes.decode("utf-8", errors="replace")

        elif ext == "pdf":
            extracted = ""
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(raw_bytes))
                text_pages = []
                for page in reader.pages:
                    text_pages.append(page.extract_text() or "")
                extracted = "\n".join(text_pages).strip()
            except Exception:
                extracted = ""

            # Fallback to OCR if standard text extraction returned empty or tiny text
            if len(extracted) < 20 and OCR_AVAILABLE:
                extracted_ocr = self._run_ocr_on_pdf_bytes(raw_bytes)
                if extracted_ocr:
                    extracted = extracted_ocr

            return extracted if extracted else None

        elif ext in ["docx", "doc"]:
            try:
                import docx
                doc = docx.Document(io.BytesIO(raw_bytes))
                full_text = []
                for para in doc.paragraphs:
                    full_text.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        full_text.append(" | ".join(cell.text.strip() for cell in row.cells))
                extracted = "\n".join(full_text).strip()
                return extracted if extracted else None
            except Exception:
                return None

        elif ext in ["xlsx", "xls"]:
            try:
                import pandas as pd
                df = pd.read_excel(io.BytesIO(raw_bytes))
                return df.to_string()
            except Exception:
                return None

        # Fallback text decode attempt
        try:
            return raw_bytes.decode("utf-8", errors="replace")
        except Exception:
            return None

    def _run_ocr_on_pdf_bytes(self, raw_bytes: bytes) -> str:
        """Helper method to run pdf2image and pytesseract on scanned document bytes."""
        ocr_text = ""
        try:
            images = convert_from_bytes(raw_bytes)
            for img in images:
                text = pytesseract.image_to_string(img)
                ocr_text += text + "\n"
        except Exception as e:
            print(f"[OCR Warning] Failed OCR processing: {e}")
        return ocr_text.strip()